#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Re-build LKlight paper #1 (engine manuscript) 4 deliverables from MD sources.

Reads 论文正稿_EN.md / 论文正稿_中文.md and produces:
  - LKlight_manuscript_EN.docx / LKlight_manuscript_CN.docx
  - LKlight_paper_EN.html     / LKlight_paper_CN.html

Figures are embedded inline from the markdown (figures/, figures3/); unlike the
paper-2 builder there is no trailing "Figures" appendix section.
"""
import os, re, sys, html
import markdown
from docx import Document
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
IMG_RE = re.compile(r"!\[([^\]]*)\]\(([^)\s]+)(?:\s+\"([^\"]*)\")?\)")
HDR_RE = re.compile(r"^(#{1,6})\s+(.*)")
TBL_RE = re.compile(r"^\|.+\|$")


def find_img(rel):
    if not rel:
        return None
    if os.path.isabs(rel) and os.path.exists(rel):
        return rel
    p = os.path.join(ROOT, rel)
    if os.path.exists(p):
        return p
    rel2 = rel.lstrip("./")
    p = os.path.join(ROOT, rel2)
    if os.path.exists(p):
        return p
    for d in ("figures", "figures2", "figures3"):
        p = os.path.join(ROOT, d, os.path.basename(rel))
        if os.path.exists(p):
            return p
    return None


def add_picture_if_any(paragraph, alt, src, width_in=6.0):
    path = find_img(src)
    if not path:
        paragraph.add_run(f"[missing image: {src}]").italic = True
        return
    run = paragraph.add_run()
    run.add_picture(path, width=Inches(width_in))


def parse_table_block(lines, start):
    rows, i = [], start
    while i < len(lines) and TBL_RE.match(lines[i]):
        rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    cleaned = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
    return cleaned, i


def build_docx(md_path, out_path, title):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    doc.add_heading(title, level=0)
    lines = open(md_path, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if TBL_RE.match(line):
            rows, ni = parse_table_block(lines, i)
            if rows:
                ncols = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncols)
                t.style = "Light List Accent 1"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        if ci >= ncols:
                            break
                        t.cell(ri, ci).text = cell
                doc.add_paragraph("")
            i = ni
            continue
        m = IMG_RE.search(line)
        if m and m.start() < 5:
            p = doc.add_paragraph()
            p.alignment = 1
            add_picture_if_any(p, m.group(1), m.group(2))
            i += 1
            continue
        m = HDR_RE.match(line)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
            i += 1
            continue
        if line.strip() == "---":
            doc.add_paragraph("─" * 40)
            i += 1
            continue
        buf = [line]
        i += 1
        while (i < len(lines) and lines[i].strip()
               and not HDR_RE.match(lines[i]) and not TBL_RE.match(lines[i])
               and not IMG_RE.search(lines[i]) and lines[i].strip() != "---"):
            buf.append(lines[i])
            i += 1
        text = " ".join(buf)
        doc.add_paragraph(text)
    doc.save(out_path)
    print("docx:", out_path)


def build_html(md_path, out_path):
    body = markdown.markdown(open(md_path, encoding="utf-8").read(),
                             extensions=["tables", "fenced_code", "sane_lists",
                                         "toc", "attr_list"],
                             output_format="html5")
    css = ("<style>body{font-family:'Times New Roman',serif;max-width:780px;"
           "margin:2em auto;line-height:1.45;padding:0 1em}"
           "table{border-collapse:collapse;margin:1em 0}"
           "th,td{border:1px solid #999;padding:4px 8px}th{background:#eee}"
           "img{max-width:100%;height:auto}h1,h2,h3{color:#0072B2}</style>")
    title = os.path.basename(md_path)
    open(out_path, "w", encoding="utf-8").write(
        f"<!doctype html><html><head><meta charset=utf-8><title>{title}</title>"
        f"{css}</head><body>\n{body}\n</body></html>")
    print("html:", out_path)


def main():
    pairs = [
        ("EN", "论文正稿_EN.md", "LKlight_manuscript_EN.docx", "LKlight_paper_EN.html",
         "LKlight: a complete, quantitatively benchmarked Rust engine for the LightDock protein docking protocol"),
        ("CN", "论文正稿_中文.md", "LKlight_manuscript_CN.docx", "LKlight_paper_CN.html",
         "LKlight：LightDock 蛋白对接协议的完整、定量基准验证的 Rust 引擎"),
    ]
    for lang, md, docx_name, html_name, title in pairs:
        md_path = os.path.join(ROOT, md)
        if not os.path.exists(md_path):
            print(f"[{lang}] skip, missing {md}")
            continue
        build_docx(md_path, os.path.join(ROOT, docx_name), title)
        build_html(md_path, os.path.join(ROOT, html_name))


if __name__ == "__main__":
    main()
