#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Re-build LKlight paper #2 (GPU/grid) 4 deliverables from the MD sources.

Reads 论文正稿2_GPU_{EN,CN}.md and produces:
  - LKlight_manuscript2_GPU_{EN,CN}.docx  (paragraphs, tables, embedded PNGs)
  - LKlight_paper2_GPU_{EN,CN}.html       (markdown -> HTML with img refs)

Image roots searched in order: figures/, figures2/, figures3/.
"""
import os, re, sys, html
import docx
from docx import Document
from docx.shared import Inches, Pt
import markdown

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))   # LKlight论文/
IMG_DIRS = [os.path.join(ROOT, "figures"),
            os.path.join(ROOT, "figures2"),
            os.path.join(ROOT, "figures3")]

def find_img(rel):
    """Resolve an image path relative to ROOT or an absolute path."""
    if not rel: return None
    if os.path.isabs(rel) and os.path.exists(rel): return rel
    p = os.path.join(ROOT, rel)
    if os.path.exists(p): return p
    # strip leading "../" or "./"
    rel2 = rel.lstrip("./")
    p = os.path.join(ROOT, rel2)
    if os.path.exists(p): return p
    # try each IMG_DIR
    for d in IMG_DIRS:
        p = os.path.join(d, os.path.basename(rel))
        if os.path.exists(p): return p
    return None

IMG_RE  = re.compile(r"!\[([^\]]*)\]\(([^)\s]+)(?:\s+\"([^\"]*)\")?\)")
HDR_RE  = re.compile(r"^(#{1,6})\s+(.*)")
TBL_RE  = re.compile(r"^\|.+\|$")

def render_md_to_html(md_text):
    return markdown.markdown(
        md_text,
        extensions=["tables", "fenced_code", "sane_lists", "toc", "attr_list"],
        output_format="html5",
    )

def add_picture_if_any(paragraph, alt, src, width_in=6.0):
    path = find_img(src)
    if not path:
        paragraph.add_run(f"[missing image: {src}]").italic = True
        return False
    try:
        run = paragraph.add_run()
        run.add_picture(path, width=Inches(width_in))
        # caption
        if alt:
            p = paragraph._p.getparent().add_paragraph()
            r = p.add_run(alt); r.italic = True; r.font.size = Pt(9)
        return True
    except Exception as e:
        paragraph.add_run(f"[image error: {e}]").italic = True
        return False

def parse_table_block(lines, start):
    """Return (rows, end_idx). rows is list[list[str]]."""
    rows=[]
    i=start
    while i < len(lines) and TBL_RE.match(lines[i]):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        rows.append(cells)
        i += 1
    # filter separator row like |---|---|---|
    cleaned = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
    return cleaned, i

FIGURES = [
    ("Figure 1 \u2014 two-scale score decomposition (panel A: LKlight; panel B: CPU grid; panel C: CUDA batch).",
     os.path.join(ROOT, "figures2/fig2p1_decomposition.png")),
    ("Figure 2 \u2014 RNA large-system per-function wall-clock (Table 1 visualisation).",
     os.path.join(ROOT, "figures2/fig2p2_speedup.png")),
    ("Figure 3 \u2014 wall-clock vs per-swarm glowworm count N (Table 3 visualisation).",
     os.path.join(ROOT, "figures2/fig2p3_scaling.png")),
    ("Figure 4 \u2014 four task-scale scenarios, three architectures (Table 4 visualisation).",
     os.path.join(ROOT, "figures2/fig2p4_scenarios.png")),
    ("Figure 5 \u2014 CPU-grid vs reference LKlight engine: per-function numerical agreement (v1.2.0, 0.5 \u00c5 grid).",
     os.path.join(ROOT, "figures3/figF1_equiv.png")),
    ("Figure 6 \u2014 CPU-grid speedup over reference LKlight engine, per scoring function.",
     os.path.join(ROOT, "figures3/figF2_speedup.png")),
]

def append_figures_to_docx(doc):
    doc.add_page_break()
    doc.add_heading("Figures", level=1)
    for caption, path in FIGURES:
        if not os.path.exists(path):
            continue
        p = doc.add_paragraph(); p.alignment = 1
        run = p.add_run(); run.add_picture(path, width=Inches(6.0))
        cap = doc.add_paragraph(); cap.alignment = 1
        r = cap.add_run(caption); r.italic = True; r.font.size = Pt(10)

def append_figures_to_html(html_text):
    extra = "\n<h1>Figures</h1>\n"
    for caption, path in FIGURES:
        if not os.path.exists(path):
            continue
        rel = os.path.relpath(path, ROOT)
        extra += f'<figure><img src="{html.escape(rel)}" alt="{html.escape(caption)}"><figcaption><em>{html.escape(caption)}</em></figcaption></figure>\n'
    return html_text + extra

def build_docx(md_path, out_path, title):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    # title
    h = doc.add_heading(title, level=0)
    lines = open(md_path, encoding="utf-8").read().splitlines()
    i = 0; pid = None
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            doc.add_paragraph(""); i += 1; continue
        # table
        if TBL_RE.match(line):
            rows, ni = parse_table_block(lines, i)
            if rows:
                ncols = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncols)
                t.style = "Light List Accent 1"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        if ci >= ncols: break
                        t.cell(ri, ci).text = cell
                doc.add_paragraph("")
            i = ni; continue
        # image
        m = IMG_RE.search(line)
        if m and m.start() < 5:   # whole-line image
            alt, src, _ = m.group(1), m.group(2), m.group(3)
            p = doc.add_paragraph(); p.alignment = 1
            add_picture_if_any(p, alt, src, width_in=6.0)
            i += 1; continue
        # heading
        m = HDR_RE.match(line)
        if m:
            level = min(len(m.group(1)), 4)
            doc.add_heading(m.group(2).strip(), level=level)
            i += 1; continue
        # horizontal rule
        if line.strip() == "---":
            doc.add_paragraph("─" * 40); i += 1; continue
        # paragraph (collect contiguous non-empty lines)
        buf = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not HDR_RE.match(lines[i]) \
              and not TBL_RE.match(lines[i]) and not IMG_RE.search(lines[i]) \
              and lines[i].strip() != "---":
            buf.append(lines[i]); i += 1
        text = " ".join(buf)
        # inline images inside paragraph? skip; render inline alt text
        text = IMG_RE.sub(lambda m: f"[{m.group(1) or 'image'}: {m.group(2)}]", text)
        doc.add_paragraph(text)
        doc.save(out_path)
        append_figures_to_docx(doc); doc.save(out_path)

def build_html(md_path, out_path):
    body = render_md_to_html(open(md_path, encoding="utf-8").read())
    body = append_figures_to_html(body)
    css = ("<style>body{font-family:'Times New Roman',serif;max-width:780px;"
           "margin:2em auto;line-height:1.45;padding:0 1em}"
           "table{border-collapse:collapse;margin:1em 0}"
           "th,td{border:1px solid #999;padding:4px 8px}"
           "th{background:#eee}"
           "img{max-width:100%;height:auto}"
           "h1,h2,h3{color:#0072B2}"
           "</style>")
    title = os.path.basename(md_path)
    open(out_path, "w", encoding="utf-8").write(
        f"<!doctype html><html><head><meta charset=utf-8><title>{title}</title>{css}</head>"
        f"<body>\n{body}\n</body></html>"
    )

def main():
    pairs = [
        ("EN", "论文正稿2_GPU_EN.md",
              "LKlight_manuscript2_GPU_EN.docx", "LKlight_paper2_GPU_EN.html",
              "Scaling LKlight to very large complexes \u2014 EN manuscript"),
        ("CN", "论文正稿2_GPU_CN.md",
              "LKlight_manuscript2_GPU_CN.docx", "LKlight_paper2_GPU_CN.html",
              "LKlight 在大复合物上的加速 \u2014 CN manuscript"),
    ]
    for lang, md, docx_name, html_name, title in pairs:
        md_path = os.path.join(ROOT, md)
        if not os.path.exists(md_path):
            print(f"[{lang}] skip, missing {md}"); continue
        docx_path = os.path.join(ROOT, docx_name)
        html_path = os.path.join(ROOT, html_name)
        print(f"[{lang}] building docx from {md} ...")
        build_docx(md_path, docx_path, title)
        print(f"[{lang}] building html from {md} ...")
        build_html(md_path, html_path)
        print(f"[{lang}] -> {docx_name} ({os.path.getsize(docx_path)} B), "
              f"{html_name} ({os.path.getsize(html_path)} B)")

if __name__ == "__main__":
    main()
